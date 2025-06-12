#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys
import json
import re
from pathlib import Path

try:
    from datafog import DataFog
    import PyPDF2
    from docx import Document
    import openpyxl
except ImportError as e:
    print(f"Biblioteca não encontrada: {e}", file=sys.stderr)
    # Continuar com regex manual se DataFog não estiver disponível

def extract_text_from_file(file_path):
    """Extrai texto de diferentes tipos de arquivo"""
    try:
        file_extension = Path(file_path).suffix.lower()
        
        if file_extension == '.pdf':
            return extract_text_from_pdf(file_path)
        elif file_extension in ['.docx', '.doc']:
            return extract_text_from_docx(file_path)
        elif file_extension in ['.xlsx', '.xls']:
            return extract_text_from_excel(file_path)
        else:
            # Arquivo de texto simples
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
    except Exception as e:
        print(f"Erro ao extrair texto: {e}", file=sys.stderr)
        return ""

def extract_text_from_pdf(file_path):
    """Extrai texto de arquivo PDF"""
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ''
            for page in pdf_reader.pages:
                text += page.extract_text() + '\n'
            return text
    except Exception as e:
        print(f"Erro ao ler PDF: {e}", file=sys.stderr)
        return ""

def extract_text_from_docx(file_path):
    """Extrai texto de arquivo DOCX"""
    try:
        doc = Document(file_path)
        text = ''
        for paragraph in doc.paragraphs:
            text += paragraph.text + '\n'
        return text
    except Exception as e:
        print(f"Erro ao ler DOCX: {e}", file=sys.stderr)
        return ""

def extract_text_from_excel(file_path):
    """Extrai texto de arquivo Excel"""
    try:
        workbook = openpyxl.load_workbook(file_path)
        text = ''
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        text += str(cell.value) + ' '
            text += '\n'
        return text
    except Exception as e:
        print(f"Erro ao ler Excel: {e}", file=sys.stderr)
        return ""

def scan_with_regex(content, enabled_patterns):
    """Escaneamento com regex brasileiro"""
    results = []
    
    brazilian_patterns = {
        'cpf': r'\b\d{3}\.\d{3}\.\d{3}-\d{2}\b',
        'cnpj': r'\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b',
        'email': r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',
        'telefone': r'\b\(?\d{2}\)?\s?\d{4,5}-?\d{4}\b',
        'cep': r'\b\d{5}-?\d{3}\b',
        'rg': r'\b\d{1,2}\.\d{3}\.\d{3}-\d{1}\b'
    }
    
    for pattern_name in enabled_patterns:
        if pattern_name in brazilian_patterns:
            matches = re.finditer(brazilian_patterns[pattern_name], content, re.IGNORECASE)
            for match in matches:
                context_start = max(0, match.start() - 30)
                context_end = min(len(content), match.end() + 30)
                context = content[context_start:context_end].strip()
                
                results.append({
                    'type': pattern_name.upper(),
                    'value': match.group(),
                    'context': context,
                    'position': match.start(),
                    'riskLevel': 'high'
                })
    
    return results

def main():
    try:
        file_path = "uploads/67de7ac5088fb0ec3b0bbe7db053ec41"
        enabled_patterns = ["cpf","cnpj","cep","rg","email"]
        
        # Extrair texto do arquivo
        content = extract_text_from_file(file_path)
        
        if not content.strip():
            print("[]")
            return
            
        # Tentar usar DataFog primeiro
        try:
            datafog = DataFog()
            scan_results = datafog.scan(content)
            
            results = []
            if hasattr(scan_results, 'matches') and scan_results.matches:
                for match in scan_results.matches:
                    results.append({
                        'type': getattr(match, 'entity_type', 'UNKNOWN').upper(),
                        'value': getattr(match, 'text', ''),
                        'context': content[max(0, getattr(match, 'start', 0)-30):getattr(match, 'end', 0)+30],
                        'position': getattr(match, 'start', 0),
                        'riskLevel': 'high'
                    })
            
            # Se DataFog não encontrou nada, usar regex
            if not results:
                results = scan_with_regex(content, enabled_patterns)
                
        except Exception as e:
            print(f"DataFog não disponível, usando regex: {e}", file=sys.stderr)
            results = scan_with_regex(content, enabled_patterns)
        
        # Imprimir resultados
        print(json.dumps(results, ensure_ascii=False, indent=2))
        
    except Exception as e:
        print(f"Erro: {e}", file=sys.stderr)
        print("[]")

if __name__ == "__main__":
    main()
